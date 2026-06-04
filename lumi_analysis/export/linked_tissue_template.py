from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


def create_linked_tissue_plot_template(
    output_path,
    treatments,
    tissues,
    n_mice,
    sheet_name="plot_input",
):
    # Create an Excel template for linked tissue plots.
    # Each treatment gets a dynamic block of columns.
    # Each column represents one tissue from one mouse.
    # The user should paste the already-computed mean values into the data area.

    output_path = Path(output_path)

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    # Header rows
    experiment_row = 1
    tissue_row = 2
    mouse_row = 3
    data_start_row = 4

    # Styling
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    spacer_fill = PatternFill("solid", fgColor="E7E6E6")
    time_fill = PatternFill("solid", fgColor="D9EAD3")

    thin_gray = Side(style="thin", color="B7B7B7")
    border = Border(
        left=thin_gray,
        right=thin_gray,
        top=thin_gray,
        bottom=thin_gray,
    )

    # Time column
    ws.cell(experiment_row, 1).value = "Time"
    ws.cell(tissue_row, 1).value = ""
    ws.cell(mouse_row, 1).value = ""

    for row in range(1, data_start_row):
        cell = ws.cell(row, 1)
        cell.fill = time_fill
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    col = 2

    for treatment_index, treatment in enumerate(treatments):
        treatment_start_col = col

        for mouse in range(1, n_mice + 1):
            for tissue in tissues:
                ws.cell(experiment_row, col).value = treatment
                ws.cell(tissue_row, col).value = tissue
                ws.cell(mouse_row, col).value = mouse

                for row in range(1, data_start_row):
                    cell = ws.cell(row, col)
                    cell.fill = header_fill
                    cell.font = Font(bold=True)
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border = border

                col += 1

        treatment_end_col = col - 1

        # Merge treatment name across its full block.
        ws.merge_cells(
            start_row=experiment_row,
            start_column=treatment_start_col,
            end_row=experiment_row,
            end_column=treatment_end_col,
        )

        # Add one empty spacer column between treatments.
        if treatment_index < len(treatments) - 1:
            for row in range(1, data_start_row):
                cell = ws.cell(row, col)
                cell.fill = spacer_fill
                cell.border = border

            ws.column_dimensions[get_column_letter(col)].width = 4
            col += 1

    # Formatting
    ws.freeze_panes = f"B{data_start_row}"

    ws.column_dimensions["A"].width = 14

    for column_index in range(2, col):
        column_letter = get_column_letter(column_index)

        if ws.column_dimensions[column_letter].width is None:
            ws.column_dimensions[column_letter].width = 14

    for row in range(1, data_start_row):
        ws.row_dimensions[row].height = 22

    # Add a note for the user.
    note_col = col + 1
    ws.cell(1, note_col).value = "Instructions"
    ws.cell(2, note_col).value = (
        "Paste Time values in column A. "
        "Paste mean signal columns under the matching treatment/tissue/mouse headers."
    )

    ws.cell(1, note_col).font = Font(bold=True)
    ws.cell(2, note_col).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions[get_column_letter(note_col)].width = 45

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    return output_path


if __name__ == "__main__":
    create_linked_tissue_plot_template(
        output_path="linked_tissue_plot_template.xlsx",
        treatments=["control F", "control M", "6% O2", "10% CO2"],
        tissues=["Liver", "Lung", "Kidney"],
        n_mice=2,
    )
