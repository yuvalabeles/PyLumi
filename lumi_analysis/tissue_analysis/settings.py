from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


def create_tissue_settings_template(
    output_path="tissue_settings.xlsx",
    sheet_name="settings",
):
    output_path = Path(output_path)

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    headers = [
        "mice",
        "tissues",
        "y max",
        "treatments",
        "check first peak after (h)",
        "output path",
    ]

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    thin_gray = Side(style="thin", color="B7B7B7")
    border = Border(
        left=thin_gray,
        right=thin_gray,
        top=thin_gray,
        bottom=thin_gray,
    )

    for col_index, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_index)
        cell.value = header
        cell.fill = header_fill
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

        ws.column_dimensions[get_column_letter(col_index)].width = 22

    ws.freeze_panes = "A2"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    return output_path


def _format_excel_value_as_text(value):
    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


def _read_non_empty_column(raw, column_name):
    if column_name not in raw.columns:
        raise ValueError(f"Missing required column: {column_name}")

    values = raw[column_name].dropna()
    values = [_format_excel_value_as_text(value) for value in values]
    values = [value for value in values if value != ""]

    return values


def _read_y_upper_limits(raw, tissues):
    if "y max" not in raw.columns:
        return None

    y_upper_limits = []

    for row_index, tissue in enumerate(tissues):
        value = raw.loc[row_index, "y max"]

        if pd.isna(value) or str(value).strip() == "":
            y_upper_limits.append(None)
        else:
            y_upper_limits.append(float(value))

    if all(value is None for value in y_upper_limits):
        return None

    return y_upper_limits


def _read_first_peak_after_by_treatment(raw, treatments):
    column_name = "check first peak after (h)"

    if column_name not in raw.columns:
        return {}

    first_peak_after_by_treatment = {}

    for row_index, treatment in enumerate(treatments):
        value = raw.loc[row_index, column_name]

        if pd.isna(value) or str(value).strip() == "":
            first_peak_after_by_treatment[treatment] = None
        else:
            first_peak_after_by_treatment[treatment] = float(value)

    return first_peak_after_by_treatment


def read_tissue_settings(
    input_path="tissue_settings.xlsx",
    sheet_name="settings",
):
    input_path = Path(input_path)

    raw = pd.read_excel(
        input_path,
        sheet_name=sheet_name,
    )

    raw.columns = [
        str(column).strip().lower()
        for column in raw.columns
    ]

    mice = _read_non_empty_column(
        raw=raw,
        column_name="mice",
    )

    tissues = _read_non_empty_column(
        raw=raw,
        column_name="tissues",
    )

    treatments = _read_non_empty_column(
        raw=raw,
        column_name="treatments",
    )

    first_peak_after_by_treatment = _read_first_peak_after_by_treatment(
        raw=raw,
        treatments=treatments,
    )

    y_upper_limits = _read_y_upper_limits(
        raw=raw,
        tissues=tissues,
    )

    output_paths = _read_non_empty_column(
        raw=raw,
        column_name="output path",
    )

    if len(mice) == 0:
        raise ValueError("No mice were found in tissue_settings.xlsx.")

    if len(tissues) == 0:
        raise ValueError("No tissues were found in tissue_settings.xlsx.")

    if len(treatments) == 0:
        raise ValueError("No treatments were found in tissue_settings.xlsx.")

    if len(output_paths) == 0:
        raise ValueError("No output path was found in tissue_settings.xlsx.")

    output_path = Path(output_paths[0])

    return mice, tissues, treatments, y_upper_limits, output_path, first_peak_after_by_treatment
