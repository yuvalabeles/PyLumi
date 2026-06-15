from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


def create_linked_tissue_plot_template(
    output_path,
    treatments,
    tissues,
    mice,
    sheet_name="plot_input",
):
    # Create an Excel template for linked tissue plots.
    # Each treatment has its own time column because experiments may have different time axes.

    output_path = Path(output_path)

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    experiment_row = 1
    tissue_row = 2
    mouse_row = 3
    data_start_row = 4

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    time_fill = PatternFill("solid", fgColor="D9EAD3")
    spacer_fill = PatternFill("solid", fgColor="E7E6E6")

    thin_gray = Side(style="thin", color="B7B7B7")
    border = Border(
        left=thin_gray,
        right=thin_gray,
        top=thin_gray,
        bottom=thin_gray,
    )

    col = 1

    for treatment_index, treatment in enumerate(treatments):
        treatment_start_col = col

        # Treatment-specific time column.
        ws.cell(experiment_row, col).value = treatment
        ws.cell(tissue_row, col).value = "Time"
        ws.cell(mouse_row, col).value = ""

        for row in range(1, data_start_row):
            cell = ws.cell(row, col)
            cell.fill = time_fill
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

        ws.column_dimensions[get_column_letter(col)].width = 14
        col += 1

        for mouse in mice:
            for tissue in tissues:
                ws.cell(experiment_row, col).value = treatment
                ws.cell(tissue_row, col).value = tissue
                ws.cell(mouse_row, col).value = mouse

                for row in range(1, data_start_row):
                    cell = ws.cell(row, col)
                    cell.fill = header_fill
                    cell.font = Font(bold=True)
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border = border

                ws.column_dimensions[get_column_letter(col)].width = 14
                col += 1

        treatment_end_col = col - 1

        ws.merge_cells(
            start_row=experiment_row,
            start_column=treatment_start_col,
            end_row=experiment_row,
            end_column=treatment_end_col,
        )

        if treatment_index < len(treatments) - 1:
            for row in range(1, data_start_row):
                cell = ws.cell(row, col)
                cell.fill = spacer_fill
                cell.border = border

            ws.column_dimensions[get_column_letter(col)].width = 4
            col += 1

    ws.freeze_panes = f"A{data_start_row}"

    for row in range(1, data_start_row):
        ws.row_dimensions[row].height = 22

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    return output_path


def read_linked_tissue_plot_template(
    input_path,
    sheet_name="plot_input",
):
    # Read a filled linked tissue plot template.
    # Each treatment block has its own Time column.

    input_path = Path(input_path)

    raw = pd.read_excel(
        input_path,
        sheet_name=sheet_name,
        header=None,
    )

    experiment_row = 0
    tissue_row = 1
    mouse_row = 2
    data_start_row = 3

    records = []

    current_experiment = None
    current_time = None

    for col in range(raw.shape[1]):
        experiment = raw.iloc[experiment_row, col]
        tissue = raw.iloc[tissue_row, col]
        mouse = raw.iloc[mouse_row, col]

        # Empty spacer column.
        if pd.isna(experiment) and pd.isna(tissue) and pd.isna(mouse):
            current_experiment = None
            current_time = None
            continue

        # Because treatment headers are merged, only the first column has the value.
        if not pd.isna(experiment):
            current_experiment = str(experiment).strip()

        if pd.isna(tissue):
            continue

        tissue = str(tissue).strip()

        # Start of a new treatment block.
        if tissue.lower() == "time":
            current_time = pd.to_numeric(
                raw.iloc[data_start_row:, col],
                errors="coerce",
            )
            continue

        if current_experiment is None or current_time is None:
            continue

        if pd.isna(mouse):
            continue

        mouse = str(mouse).strip()

        values = pd.to_numeric(
            raw.iloc[data_start_row:, col],
            errors="coerce",
        )

        sample_df = pd.DataFrame({
            "time": current_time,
            "experiment": current_experiment,
            "tissue": tissue,
            "mouse": mouse,
            "mean": values,
        })

        sample_df = sample_df.dropna(subset=["time", "mean"])
        records.append(sample_df)

    if not records:
        raise ValueError("No valid data columns were found in the template.")

    return pd.concat(records, ignore_index=True)
