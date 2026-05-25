from pathlib import Path
from ast import literal_eval

from openpyxl import load_workbook


GROUP_OVERRIDE_SETTINGS = [
    "y_limit",
    "description",
    "peak_txt_dy",
    "mean_replicate_cols",
    "visible_replicate_cols",
]


def parse_setting_value(value):
    """
    Convert Excel values into proper Python values.
    """

    if value is None:
        return None

    if isinstance(value, (int, float, bool)):
        return value

    if isinstance(value, str):
        value = value.strip()

        if value == "":
            return None

        if value.lower() == "none":
            return None

        try:
            return literal_eval(value)

        except (ValueError, SyntaxError):
            return value

    return value


def parse_override_list(value):
    """
    Parse override rows like:

    [500, , 700]

    into:

    [500, None, 700]
    """

    if value is None:
        return []

    if not isinstance(value, str):
        return value

    text = value.strip()

    if text == "":
        return []

    # Replace empty entries with None
    text = text.replace(", ,", ", None,")
    text = text.replace(",  ,", ", None,")
    text = text.replace("[ ,", "[None,")
    text = text.replace(", ]", ", None]")

    try:
        parsed = literal_eval(text)

        if isinstance(parsed, list):
            return parsed

    except (ValueError, SyntaxError):
        pass

    return []


def load_settings(settings_path):
    settings_path = Path(settings_path)

    workbook = load_workbook(settings_path, data_only=True)
    worksheet = workbook["Settings"]

    config = {}

    # Current template columns:
    # A = Section
    # B = Setting
    # C = Value
    # D = Possible Values
    # E = Description
    # F = Config Key
    VALUE_COL = 3
    CONFIG_KEY_COL = 6

    override_lists = {}

    for row in range(2, worksheet.max_row + 1):

        config_key = worksheet.cell(
            row=row,
            column=CONFIG_KEY_COL
        ).value

        if config_key is None:
            continue

        config_key = str(config_key).strip()

        raw_value = worksheet.cell(
            row=row,
            column=VALUE_COL
        ).value

        # GROUP NAMES
        if config_key == "group_names":
            override_lists[config_key] = parse_override_list(raw_value)
            continue

        # REGULAR SETTINGS
        config[config_key] = parse_setting_value(raw_value)

        # GROUP OVERRIDE SETTINGS
        if config_key in GROUP_OVERRIDE_SETTINGS:
            override_lists[config_key] = parse_override_list(raw_value)

    # ------------------------------------------------------------------
    # BUILD plot_group_settings
    # ------------------------------------------------------------------

    group_names = override_lists.get("group_names", [])

    plot_group_settings = {}

    for group_index, group_name in enumerate(group_names):

        if group_name is None:
            continue

        group_name = str(group_name)

        group_overrides = {}

        for setting_name in GROUP_OVERRIDE_SETTINGS:

            values_list = override_lists.get(setting_name, [])

            if group_index >= len(values_list):
                continue

            value = values_list[group_index]

            # None means "no override"
            if value is None:
                continue

            group_overrides[setting_name] = value

        if len(group_overrides) > 0:
            plot_group_settings[group_name] = group_overrides

    config["plot_group_settings"] = plot_group_settings

    return config
